"""Real bounded subprocesses, document parsers and PermissionEngine; no user files or cloud."""

import asyncio
import json
import sys
from contextlib import closing
from pathlib import Path
from threading import Event
from typing import Any

import pytest
from tests.unit.test_filesystem import policy as policy

from jarvis.core.planner.contracts import PlannerInput, Proposal
from jarvis.core.planner.runner import Runner
from jarvis.observability.audit import AuditLog
from jarvis.permissions.approvals import Action, ApprovalStore
from jarvis.permissions.engine import Outcome, PermissionEngine
from jarvis.permissions.policies import Mode, Risk, Status
from jarvis.platforms.filesystem.transport import FileWorker
from jarvis.security.filesystem_policy import FilePolicy
from jarvis.tools.base import ExecutionContext, ToolError
from jarvis.tools.filesystem import DocumentResult, ReadDocument, SandboxEval, register_filesystem
from jarvis.tools.registry import ToolRegistry


def pdf_fixture(path: Path) -> None:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 30 700 Td (Contract amount: 1250.50 USD) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    writer.write(path)


@pytest.mark.asyncio
async def test_real_worker_reads_pdf_and_docx_tables(policy: FilePolicy) -> None:
    from docx import Document

    root = policy.path("Desktop")
    pdf_fixture(root / "contract.pdf")
    document = Document()
    document.add_paragraph("Договор №1")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Сумма", "1250.50"
    document.save(str(root / "contract.docx"))
    for name in ("contract.pdf", "contract.docx"):
        result = await FileWorker(policy).call(
            "fs.read_document", ReadDocument(path=str(root / name)), ExecutionContext(Event())
        )
        assert isinstance(result, DocumentResult)
        assert "1250.50" in result.text and result.version.sha256


@pytest.mark.asyncio
async def test_pdf_contract_to_csv_runs_without_gui_or_cloud(
    policy: FilePolicy, tmp_path: Path
) -> None:
    pdf_fixture(policy.path("Desktop") / "contract.pdf")
    registry = ToolRegistry()
    register_filesystem(registry, policy)
    with closing(AuditLog(tmp_path / "audit.sqlite3")) as audit:
        store = ApprovalStore()
        engine = PermissionEngine(registry, store, audit)
        authority = store.take_authority(audit.approved)

        class Provider:
            async def propose(self, data: PlannerInput) -> Proposal:
                def call(tool: str, **args: object) -> Proposal:
                    return Proposal(
                        kind="call", tool=tool, arguments=json.dumps(args, ensure_ascii=False)
                    )

                if not data.steps:
                    return call("fs.search", query=".pdf", directory="Desktop")
                last = json.loads(data.steps[-1].outcome.result_json or "{}")
                if len(data.steps) == 1:
                    return call("fs.read_document", path=last["files"][0]["path"])
                if len(data.steps) == 2:
                    return call(
                        "python.sandbox_eval",
                        code="amount = float(data.split('amount: ')[1].split()[0])\n"
                        "csv.format([{'amount': amount, 'currency': 'USD'}])",
                        data=last["text"],
                    )
                if len(data.steps) == 3:
                    return call(
                        "fs.write_file",
                        path="workspace/contract.csv",
                        content=json.loads(last["value_json"]),
                    )
                return Proposal(kind="finish")

        async def approve(action: Action) -> Any:
            return authority.approve(action)

        async def clarify(question: str) -> str | None:
            raise AssertionError("No ambiguity in fixture")

        result = await Runner(registry, engine, Provider(), approve, clarify).run(
            "Прочитай PDF-договор и занеси сумму в CSV", Mode.EXECUTE
        )
        assert result.status == "finished", result
        assert len(result.steps) == 4 and all(
            step.outcome.status is Status.SUCCESS for step in result.steps
        )
        assert (
            Path(policy.workspace) / "contract.csv"
        ).read_text() == "amount,currency\n1250.5,USD\n"
        records = "\n".join(audit.recent())
        assert (
            "1250.5" not in records
            and "Contract amount" not in records
            and policy.workspace not in records
        )


@pytest.mark.asyncio
async def test_create_safe_overwrite_confirm_and_version_change(
    policy: FilePolicy, tmp_path: Path
) -> None:
    registry = ToolRegistry()
    register_filesystem(registry, policy)
    with closing(AuditLog(tmp_path / "audit.sqlite3")) as audit:
        store = ApprovalStore()
        engine = PermissionEngine(registry, store, audit)
        authority = store.take_authority(audit.approved)

        def prepare(tool: str, **args: object) -> Action:
            action = engine.prepare(tool, args, Mode.EXECUTE)
            assert isinstance(action, Action), action
            return action

        created = prepare("fs.write_file", path="workspace/file.csv", content="old")
        assert created.risk is Risk.SAFE
        assert (await engine.execute(created)).status is Status.SUCCESS
        path = str(Path(policy.workspace) / "file.csv")
        read = await engine.execute(prepare("fs.read_document", path=path))
        version = json.loads(read.result_json or "{}")["version"]
        denied = prepare("fs.write_file", path=path, content="new", expected_version=version)
        assert denied.risk is Risk.CONFIRM
        assert (await engine.execute(denied)).status is Status.DENIED
        assert Path(path).read_text() == "old"
        changed = prepare("fs.write_file", path=path, content="new", expected_version=version)
        token = authority.approve(changed)
        Path(path).write_text("user edit")
        assert (await engine.execute(changed, token)).status is Status.ERROR
        assert Path(path).read_text() == "user edit"
        refreshed = await engine.execute(prepare("fs.read_document", path=path))
        version = json.loads(refreshed.result_json or "{}")["version"]
        confirmed = prepare(
            "fs.write_file", path=path, content="approved", expected_version=version
        )
        token = authority.approve(confirmed)
        assert (await engine.execute(confirmed, token)).status is Status.SUCCESS
        assert Path(path).read_text() == "approved"
        assert (await engine.execute(confirmed, token)).status is Status.INVALID


@pytest.mark.asyncio
async def test_simulation_never_calls_worker_or_touches_disk(
    policy: FilePolicy, tmp_path: Path
) -> None:
    class NoWorker:
        async def call(self, *args: Any) -> Any:
            raise AssertionError("Simulation must never invoke an adapter")

    registry = ToolRegistry()
    register_filesystem(registry, policy, NoWorker())
    with closing(AuditLog(tmp_path / "audit.sqlite3")) as audit:
        engine = PermissionEngine(registry, ApprovalStore(), audit)
        for name, args in (
            ("fs.search", {"query": ".pdf", "directory": "Desktop"}),
            ("fs.read_document", {"path": "Desktop/missing.pdf"}),
            ("fs.write_file", {"path": "workspace/new.csv", "content": "x"}),
            ("python.sandbox_eval", {"code": "sum(range(10))"}),
        ):
            action = engine.prepare(name, args, Mode.SIMULATION)
            assert isinstance(action, Action), action
            assert (await engine.execute(action)).status is Status.SIMULATED
        assert not Path(policy.workspace).exists()
        rejected = engine.prepare("python.sandbox_eval", {"code": "import os"})
        assert isinstance(rejected, Outcome) and rejected.status is Status.INVALID


@pytest.mark.asyncio
@pytest.mark.parametrize("cancel", [True, False])
async def test_hung_worker_is_killed_reaped_and_receives_no_credentials(
    policy: FilePolicy, monkeypatch: pytest.MonkeyPatch, cancel: bool
) -> None:
    launch = asyncio.create_subprocess_exec
    processes: list[asyncio.subprocess.Process] = []

    async def spawn(*args: Any, **kwargs: Any) -> asyncio.subprocess.Process:
        assert "-I" in args and "-B" in args
        assert "OPENAI_API_KEY" not in kwargs["env"] and "PYTHONPATH" not in kwargs["env"]
        process = await launch(sys.executable, "-I", "-c", "import time;time.sleep(60)", **kwargs)
        processes.append(process)
        return process

    monkeypatch.setenv("OPENAI_API_KEY", "SYNTHETIC_DO_NOT_FORWARD")
    monkeypatch.setattr(asyncio, "create_subprocess_exec", spawn)
    task = asyncio.create_task(
        FileWorker(policy, timeout=0.15).call(
            "python.sandbox_eval", SandboxEval(code="1+1"), ExecutionContext(Event())
        )
    )
    while not processes:
        await asyncio.sleep(0.005)
    if cancel:
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
    else:
        with pytest.raises(ToolError, match="worker_timeout"):
            await task
    assert processes[0].returncode is not None


@pytest.mark.asyncio
async def test_unsupported_corrupt_and_image_only_documents_are_explicit(
    policy: FilePolicy,
) -> None:
    from pypdf import PdfWriter

    root = policy.path("Desktop")
    worker = FileWorker(policy)
    for name in ("broken.pdf", "broken.docx"):
        (root / name).write_bytes(b"not a document")
        with pytest.raises(ToolError):
            await worker.call(
                "fs.read_document", ReadDocument(path=str(root / name)), ExecutionContext(Event())
            )
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(root / "no-text.pdf")
    result = await worker.call(
        "fs.read_document", ReadDocument(path=str(root / "no-text.pdf")), ExecutionContext(Event())
    )
    assert isinstance(result, DocumentResult) and result.notice == "no_text_layer"
    assert not result.text.strip()


@pytest.mark.asyncio
async def test_runner_rejects_unobserved_overwrite_before_approval(
    policy: FilePolicy, tmp_path: Path
) -> None:
    registry = ToolRegistry()
    register_filesystem(registry, policy)
    from jarvis.platforms.filesystem.documents import read_document
    from jarvis.platforms.filesystem.storage import write
    from jarvis.tools.filesystem import WriteFile

    created = write(policy, WriteFile(path="workspace/report.txt", content="keep"))
    observed = read_document(policy, ReadDocument(path=created.path))

    class Provider:
        async def propose(self, data: PlannerInput) -> Proposal:
            return Proposal(
                kind="call",
                tool="fs.write_file",
                arguments=json.dumps(
                    {
                        "path": observed.path,
                        "content": "bad",
                        "expected_version": observed.version.model_dump(),
                    }
                ),
            )

    async def approve(action: Action) -> Any:
        raise AssertionError("An unobserved version must not reach approval")

    async def clarify(question: str) -> str | None:
        return None

    with closing(AuditLog(tmp_path / "audit.sqlite3")) as audit:
        engine = PermissionEngine(registry, ApprovalStore(), audit)
        result = await Runner(registry, engine, Provider(), approve, clarify).run(
            "Обнови файл", Mode.EXECUTE
        )
        assert result.error == "unobserved_target"
        assert Path(observed.path).read_text() == "keep"
